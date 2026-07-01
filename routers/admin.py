import csv
import hashlib
import io
import os
import re
import threading
import uuid
from datetime import datetime, timezone
from typing import Optional

import openpyxl

import bcrypt
from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from slowapi import Limiter
from slowapi.util import get_remote_address
from sqlalchemy import func, desc as sqldesc
from sqlalchemy.orm import Session

import models
import schemas as _schemas
from database import SessionLocal, get_db
from services.images import read_limited, compress_image_async
from cache import _hospitales_cache, cache_clear

BASE_URL = os.getenv("BASE_URL", "https://rescate.ventatalk.com")
UPLOAD_DIR = "uploads/capturas"

limiter = Limiter(key_func=get_remote_address)
router = APIRouter(prefix="/admin", tags=["admin"])
templates = Jinja2Templates(directory="templates")

COOKIE_NAME = "rescate_session"
SESSION_MAX_AGE = 8 * 3600  # 8 horas en segundos

# Job store en memoria (se limpia al reiniciar el container)
_jobs: dict = {}
# Estructura de cada job:
# { "status": "pending|running|done|error",
#   "progress": 0-100,
#   "result": None | {...},
#   "error": None | str,
#   "created_at": datetime }

def get_serializer():
    secret = os.getenv("SECRET_KEY", "change-me")
    return URLSafeTimedSerializer(secret)

def generate_csrf_token(session_data: str) -> str:
    secret = os.getenv("SECRET_KEY", "")
    raw = f"{session_data}{secret}csrf"
    return hashlib.sha256(raw.encode()).hexdigest()[:32]

def validate_csrf(request: Request, token: str, session_data: str) -> bool:
    expected = generate_csrf_token(session_data)
    return token == expected

def get_current_admin(request: Request):
    token = request.cookies.get(COOKIE_NAME)
    if not token:
        return None
    try:
        data = get_serializer().loads(token, max_age=SESSION_MAX_AGE)
        return data.get("username")
    except (SignatureExpired, BadSignature):
        return None

def require_admin(request: Request):
    user = get_current_admin(request)
    if not user:
        raise RedirectToLogin()
    return user

class RedirectToLogin(Exception):
    pass

# --- Login ---

@router.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    csrf_token = generate_csrf_token("login")
    return templates.TemplateResponse("admin/login.html", {"request": request, "error": None, "csrf_token": csrf_token})

@router.post("/login")
@limiter.limit("5/15minutes")
async def login_post(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    if not validate_csrf(request, csrf_token, "login"):
        return templates.TemplateResponse(
            "admin/login.html",
            {"request": request, "error": "Token de seguridad inválido. Recarga la página.", "csrf_token": generate_csrf_token("login")},
            status_code=403,
        )
    user = db.query(models.AdminUser).filter(models.AdminUser.username == username).first()
    if user and bcrypt.checkpw(password.encode(), user.password_hash.encode()):
        token = get_serializer().dumps({"username": username})
        response = RedirectResponse(url="/admin/dashboard", status_code=303)
        response.set_cookie(
            COOKIE_NAME, token,
            max_age=SESSION_MAX_AGE,
            httponly=True,
            secure=False,  # True en producción con HTTPS
            samesite="lax",
        )
        return response
    return templates.TemplateResponse(
        "admin/login.html",
        {"request": request, "error": "Usuario o contraseña incorrectos", "csrf_token": generate_csrf_token("login")},
        status_code=401,
    )

@router.get("/logout")
async def logout():
    response = RedirectResponse(url="/admin/login", status_code=303)
    response.delete_cookie(COOKIE_NAME)
    return response

# --- Dashboard ---

@router.get("/dashboard", response_class=HTMLResponse)
async def dashboard(request: Request, db: Session = Depends(get_db)):
    admin = get_current_admin(request)
    if not admin:
        return RedirectResponse(url="/admin/login", status_code=303)

    personas = db.query(models.PersonaDesaparecida).order_by(
        models.PersonaDesaparecida.created_at.desc()
    ).all()

    hoy = datetime.now(timezone.utc).date()
    total = len(personas)
    hoy_count = sum(1 for p in personas if p.created_at and p.created_at.date() == hoy)
    con_foto = sum(1 for p in personas if p.foto_url)
    encontrados = sum(1 for p in personas if p.estado == "encontrado")
    en_proceso = sum(1 for p in personas if p.estado == "en_proceso")

    csrf_token = generate_csrf_token(admin)
    return templates.TemplateResponse("admin/dashboard.html", {
        "request": request,
        "admin": admin,
        "personas": personas,
        "total": total,
        "hoy": hoy_count,
        "con_foto": con_foto,
        "encontrados": encontrados,
        "en_proceso": en_proceso,
        "csrf_token": csrf_token,
    })

# --- Cambiar estado ---

@router.post("/personas/{persona_id}/estado")
async def cambiar_estado(
    persona_id: int,
    request: Request,
    estado: str = Form(...),
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        return RedirectResponse(url="/admin/login", status_code=303)

    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")

    if estado not in ("desaparecido", "encontrado", "en_proceso", "localizado"):
        return RedirectResponse(url="/admin/dashboard", status_code=303)

    persona = db.query(models.PersonaDesaparecida).filter(
        models.PersonaDesaparecida.id == persona_id
    ).first()
    if persona:
        persona.estado = estado
        db.commit()
    return RedirectResponse(url="/admin/dashboard", status_code=303)

# --- Editar persona (admin) ---

@router.patch("/personas/{persona_id}")
async def editar_persona_admin(
    persona_id: int,
    request: Request,
    nombres_apellidos: Optional[str] = Form(None),
    cedula: Optional[str] = Form(None),
    ultima_ubicacion: Optional[str] = Form(None),
    descripcion: Optional[str] = Form(None),
    numero_contacto: Optional[str] = Form(None),
    numero_contacto_2: Optional[str] = Form(None),
    quien_ayudo: Optional[str] = Form(None),
    contacto_quien_ayudo: Optional[str] = Form(None),
    estado: Optional[str] = Form(None),
    tipo_reporte: Optional[str] = Form(None),
    estado_clinico: Optional[str] = Form(None),
    sexo: Optional[str] = Form(None),
    edad_aproximada: Optional[str] = Form(None),
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")

    persona = db.query(models.PersonaDesaparecida).filter(
        models.PersonaDesaparecida.id == persona_id
    ).first()
    if not persona:
        raise HTTPException(status_code=404, detail="No encontrado")

    if cedula is not None:
        cedula = re.sub(r'[^0-9]', '', cedula.strip()) or None

    ESTADOS_VALIDOS = {"desaparecido", "encontrado", "en_proceso", "localizado"}
    TIPOS_VALIDOS = {"desaparecido", "encontrado_vivo"}
    CLINICO_VALIDOS = {"atrapado", "herido", "inconsciente", "fallecido", "sin_informacion"}
    SEXO_VALIDOS = {"masculino", "femenino", "no_determinado"}
    EDAD_VALIDOS = {"nino", "joven", "adulto", "adulto_mayor", "no_sabe"}

    if nombres_apellidos and nombres_apellidos.strip():
        persona.nombres_apellidos = nombres_apellidos.strip()
    if cedula is not None:
        persona.cedula = cedula
    if ultima_ubicacion and ultima_ubicacion.strip():
        persona.ultima_ubicacion = ultima_ubicacion.strip()
    if descripcion is not None:
        persona.descripcion = descripcion.strip() or None
    if numero_contacto and numero_contacto.strip():
        persona.numero_contacto = numero_contacto.strip()
    if numero_contacto_2 is not None:
        persona.numero_contacto_2 = numero_contacto_2.strip() or None
    if quien_ayudo is not None:
        persona.quien_ayudo = quien_ayudo.strip() or None
    if contacto_quien_ayudo is not None:
        persona.contacto_quien_ayudo = contacto_quien_ayudo.strip() or None
    if estado and estado in ESTADOS_VALIDOS:
        persona.estado = estado
    if tipo_reporte and tipo_reporte in TIPOS_VALIDOS:
        persona.tipo_reporte = tipo_reporte
    if estado_clinico and estado_clinico in CLINICO_VALIDOS:
        persona.estado_clinico = estado_clinico
    if sexo and sexo in SEXO_VALIDOS:
        persona.sexo = sexo
    if edad_aproximada and edad_aproximada in EDAD_VALIDOS:
        persona.edad_aproximada = edad_aproximada

    db.commit()
    return {"ok": True, "id": persona_id}


# --- Eliminar persona ---

@router.delete("/personas/{persona_id}")
async def eliminar_persona(
    persona_id: int,
    request: Request,
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")

    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")

    persona = db.query(models.PersonaDesaparecida).filter(
        models.PersonaDesaparecida.id == persona_id
    ).first()
    if not persona:
        raise HTTPException(status_code=404, detail="No encontrado")
    db.delete(persona)
    db.commit()
    return {"ok": True}

# --- Eliminar bombero ---

@router.delete("/bomberos/{reporte_id}")
async def eliminar_bombero(
    reporte_id: int,
    request: Request,
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")

    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")

    reporte = db.query(models.BomberoReporte).filter(
        models.BomberoReporte.id == reporte_id
    ).first()
    if not reporte:
        raise HTTPException(status_code=404, detail="No encontrado")
    db.delete(reporte)
    db.commit()
    return {"ok": True}

# --- Export CSV ---

@router.get("/export/csv")
async def export_csv(request: Request, db: Session = Depends(get_db)):
    admin = get_current_admin(request)
    if not admin:
        return RedirectResponse(url="/admin/login", status_code=303)

    personas = db.query(models.PersonaDesaparecida).order_by(
        models.PersonaDesaparecida.created_at.desc()
    ).all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "ID", "Nombres y Apellidos", "Cédula", "Última Ubicación",
        "Foto URL", "Descripción", "Número Contacto", "Número Contacto 2",
        "Quién Ayudó", "Contacto Quien Ayudó", "Estado", "Fecha Registro"
    ])
    for p in personas:
        writer.writerow([
            p.id, p.nombres_apellidos, p.cedula or "", p.ultima_ubicacion,
            p.foto_url or "", p.descripcion or "", p.numero_contacto,
            p.numero_contacto_2 or "",
            p.quien_ayudo or "", p.contacto_quien_ayudo or "",
            p.estado, p.created_at.strftime("%Y-%m-%d %H:%M:%S") if p.created_at else "",
        ])

    output.seek(0)
    fecha = datetime.now().strftime("%Y%m%d_%H%M%S")
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=rescate_{fecha}.csv"},
    )

# --- Importar CSV ---

_VALID_ESTADOS_PAC = {"ingresado", "estable", "grave", "critico", "alta", "trasladado", "fallecido"}


@router.post("/import/pacientes")
async def import_pacientes(
    request: Request,
    file: UploadFile = File(...),
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")
    if not (file.filename or "").lower().endswith(".csv"):
        raise HTTPException(status_code=400, detail="El archivo debe ser .csv")

    content = await file.read()
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")

    reader = csv.DictReader(io.StringIO(text))
    importados = 0
    errores = 0
    detalle: list[str] = []

    for i, row in enumerate(reader, start=2):
        row_n = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        nombres = row_n.get("nombres_apellidos", "")
        if not nombres:
            continue
        try:
            hosp_id_str = row_n.get("hospital_id", "")
            hosp_id = int(hosp_id_str) if hosp_id_str else None
            estado = row_n.get("estado_paciente", "").lower() or "ingresado"
            if estado not in _VALID_ESTADOS_PAC:
                estado = "ingresado"
            paciente = models.PacienteHospitalizado(
                nombres_apellidos=nombres,
                nombre_hospital=row_n.get("nombre_hospital") or None,
                hospital_id=hosp_id,
                edad=row_n.get("edad") or None,
                sexo=row_n.get("sexo") or None,
                cedula=row_n.get("cedula") or None,
                parentesco=row_n.get("parentesco") or None,
                procedencia=row_n.get("procedencia") or None,
                observaciones=row_n.get("observaciones") or None,
                datos_adicionales=row_n.get("datos_adicionales") or None,
                estado_paciente=estado,
                reportado_por=row_n.get("reportado_por") or None,
            )
            db.add(paciente)
            importados += 1
        except Exception as exc:
            errores += 1
            detalle.append(f"Fila {i}: {exc}")

    db.commit()
    return {"importados": importados, "errores": errores, "detalle": detalle}


@router.post("/import/hospitales")
async def import_hospitales(
    request: Request,
    file: UploadFile = File(...),
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")
    if not (file.filename or "").lower().endswith(".csv"):
        raise HTTPException(status_code=400, detail="El archivo debe ser .csv")

    content = await file.read()
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")

    reader = csv.DictReader(io.StringIO(text))
    importados = 0
    errores = 0
    detalle: list[str] = []

    for i, row in enumerate(reader, start=2):
        row_n = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        nombre = row_n.get("nombre", "")
        zona = row_n.get("zona", "")
        if not nombre or not zona:
            if nombre or zona:
                errores += 1
                detalle.append(f"Fila {i}: faltan nombre o zona requeridos")
            continue
        try:
            activo_str = row_n.get("activo", "").lower()
            activo = activo_str not in ("false", "0", "no")
            hospital = models.Hospital(
                nombre=nombre,
                zona=zona,
                direccion=row_n.get("direccion") or None,
                telefono=row_n.get("telefono") or None,
                tipo=row_n.get("tipo") or None,
                activo=activo,
            )
            db.add(hospital)
            importados += 1
        except Exception as exc:
            errores += 1
            detalle.append(f"Fila {i}: {exc}")

    db.commit()
    return {"importados": importados, "errores": errores, "detalle": detalle}


# --- Nuevo hospital ---

@router.post("/hospitales/nuevo")
async def nuevo_hospital(
    request: Request,
    nombre: str = Form(...),
    zona: str = Form(...),
    direccion: Optional[str] = Form(None),
    telefono: Optional[str] = Form(None),
    tipo: Optional[str] = Form(None),
    activo: Optional[str] = Form(None),
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")
    hospital = models.Hospital(
        nombre=nombre,
        zona=zona,
        direccion=direccion or None,
        telefono=telefono or None,
        tipo=tipo or None,
        activo=activo is not None,
    )
    db.add(hospital)
    db.commit()
    cache_clear(_hospitales_cache)
    db.refresh(hospital)
    return {"ok": True, "id": hospital.id, "nombre": hospital.nombre}


# --- Nuevo paciente (admin) ---

@router.post("/pacientes/nuevo")
async def nuevo_paciente_admin(
    request: Request,
    nombres_apellidos: str = Form(...),
    hospital_id: Optional[int] = Form(None),
    nombre_hospital: Optional[str] = Form(None),
    edad: Optional[str] = Form(None),
    sexo: Optional[str] = Form(None),
    cedula: Optional[str] = Form(None),
    parentesco: Optional[str] = Form(None),
    procedencia: Optional[str] = Form(None),
    observaciones: Optional[str] = Form(None),
    datos_adicionales: Optional[str] = Form(None),
    estado_paciente: str = Form("ingresado"),
    reportado_por: Optional[str] = Form(None),
    necesita_ayuda: bool = Form(False),
    tipo_ayuda: Optional[str] = Form(None),
    foto_captura: UploadFile = File(None),
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")

    foto_url = None
    if foto_captura and foto_captura.filename:
        file_bytes = await read_limited(foto_captura)
        if file_bytes is None:
            raise HTTPException(status_code=413, detail="Archivo demasiado grande (máx 10MB)")
        compressed = await compress_image_async(file_bytes)
        if compressed:
            filename = f"cap_{uuid.uuid4()}.jpg"
            with open(os.path.join(UPLOAD_DIR, filename), "wb") as f:
                f.write(compressed)
            foto_url = f"{BASE_URL}/uploads/capturas/{filename}"

    paciente = models.PacienteHospitalizado(
        hospital_id=hospital_id,
        nombre_hospital=nombre_hospital or None,
        nombres_apellidos=nombres_apellidos,
        edad=edad or None,
        sexo=sexo or None,
        cedula=cedula or None,
        parentesco=parentesco or None,
        procedencia=procedencia or None,
        observaciones=observaciones or None,
        datos_adicionales=datos_adicionales or None,
        estado_paciente=estado_paciente or "ingresado",
        reportado_por=reportado_por or None,
        necesita_ayuda=necesita_ayuda,
        tipo_ayuda=tipo_ayuda or None,
        foto_captura_url=foto_url,
    )
    db.add(paciente)
    db.commit()
    db.refresh(paciente)
    return {"ok": True, "id": paciente.id}


# --- Eliminar todos los pacientes ---

@router.delete("/pacientes/todos")
async def eliminar_todos_pacientes(
    request: Request,
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")
    count = db.query(models.PacienteHospitalizado).count()
    db.query(models.PacienteHospitalizado).delete()
    db.commit()
    return {"ok": True, "eliminados": count}


# --- Eliminar paciente por id ---

@router.delete("/pacientes/{paciente_id}")
async def eliminar_paciente(
    paciente_id: int,
    request: Request,
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")
    paciente = db.query(models.PacienteHospitalizado).filter(
        models.PacienteHospitalizado.id == paciente_id
    ).first()
    if not paciente:
        raise HTTPException(status_code=404, detail="No encontrado")
    db.delete(paciente)
    db.commit()
    return {"ok": True}


# --- Helpers para importación Excel ---

def parse_estado_from_observaciones(obs: str) -> str:
    if "ESTADO EN CONFLICTO" in obs.upper():
        return "ingresado"
    obs_lower = obs.lower()
    if "fallecid" in obs_lower:
        return "fallecido"
    if re.search(r'\balta\b', obs_lower):
        return "alta"
    if "traslad" in obs_lower:
        return "trasladado"
    if "grave" in obs_lower:
        return "grave"
    if "critico" in obs_lower or "crítico" in obs_lower:
        return "critico"
    if "estable" in obs_lower:
        return "estable"
    if "internado" in obs_lower or "ingresado" in obs_lower:
        return "ingresado"
    return "ingresado"


def limpiar_numero_excel(valor) -> str:
    if valor is None:
        return None
    s = str(valor).strip()
    if s in ('', 'None', 'nan', '\xa0'):
        return None
    if s.endswith('.0'):
        s = s[:-2]
    return s


def normalizar_cedula(valor) -> str:
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    v = str(valor).strip().upper()
    if v.endswith('.0'):
        v = v[:-2]
    return v


HOSPITAL_ALIASES = {
    "Periférico de Catia": "Hospital Periférico de Catia",
    "H. Jose Maria Vargas La Guaira": "Hospital José María Vargas (La Guaira)",
    "Hospital Militar Universitario Dr. Carlos Arvelo": "Hospital Militar Dr. Carlos Arvelo",
    "Hospital Ana Francisca Pérez de": "Hospital Ana Francisca Pérez de León",
    "Hospital Dr. José Gregorio Hernández": "Hospital Dr. José Gregorio Hernández",
    "Hospital Ricardo Baquero González": "Hospital Ricardo Baquero González",
    "Otros / Sin Clasificar": "Otros - Sin Clasificar",
    "Parque Alí Primera (Parque del Oeste)": "Parque Alí Primera-Oeste",
    "Rescatados — Alcaldía de Chacao": "Rescatados Alcaldía Chacao",
    "IVSS Hospital General Estadal de Misiones": "Hosp. Estadal Misiones IVSS",
    "Hospital General de Lídice (Dr. Jesús Yerena)": "H. General de Lidice",
}


def normalizar_hospital(nombre: str) -> str:
    if not nombre:
        return nombre
    nombre = nombre.strip()
    return HOSPITAL_ALIASES.get(nombre, nombre)


# --- Importar Excel de pacientes ---

_EXCEL_SHEET = "🔍 BUSCAR PACIENTES"


def _procesar_excel(job_id: str, file_bytes: bytes, batch_date_str: str, username: str):
    db = SessionLocal()
    try:
        _jobs[job_id]["status"] = "running"

        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as exc:
            _jobs[job_id].update({"status": "error", "error": f"No se pudo leer el archivo Excel: {exc}"})
            return

        if _EXCEL_SHEET not in wb.sheetnames:
            _jobs[job_id].update({"status": "error", "error": f"No se encontró la hoja '{_EXCEL_SHEET}'"})
            return

        ws = wb[_EXCEL_SHEET]

        def _str(val):
            if val is None:
                return None
            if isinstance(val, float) and val.is_integer():
                return str(int(val))
            s = str(val).strip()
            return s if s else None

        # Recopilar filas con datos para poder calcular progreso
        data_rows = []
        for row_num, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
            nombre_raw = row[2] if len(row) > 2 else None
            if nombre_raw and str(nombre_raw).strip():
                data_rows.append((row_num, row))

        total = len(data_rows)
        nuevos = 0
        actualizados = 0
        movimientos = 0
        errores = 0
        errores_detalle: list[str] = []

        for i, (row_num, row) in enumerate(data_rows, start=1):
            nombre_raw = str(row[2]).strip()

            try:
                if "/" in nombre_raw:
                    idx = nombre_raw.index("/")
                    nombre_principal = nombre_raw[:idx].strip()
                    variantes = nombre_raw[idx + 1:].strip() or None
                else:
                    nombre_principal = nombre_raw
                    variantes = None

                if not nombre_principal:
                    continue

                hospital = normalizar_hospital(_str(row[1] if len(row) > 1 else None))
                edad = limpiar_numero_excel(row[3] if len(row) > 3 else None)
                cedula_raw = limpiar_numero_excel(row[4] if len(row) > 4 else None)
                cedula = normalizar_cedula(cedula_raw) if cedula_raw is not None else None
                if cedula == "":
                    cedula = None
                telefono = limpiar_numero_excel(row[5] if len(row) > 5 else None)
                direccion = _str(row[6] if len(row) > 6 else None)
                observaciones = _str(row[7] if len(row) > 7 else None)
                estado = parse_estado_from_observaciones(observaciones or "")

                paciente = None
                if cedula:
                    paciente = db.query(models.PacienteHospitalizado).filter(
                        models.PacienteHospitalizado.cedula == cedula
                    ).first()

                if paciente is None:
                    paciente = db.query(models.PacienteHospitalizado).filter(
                        func.lower(models.PacienteHospitalizado.nombres_apellidos) == nombre_principal.lower()
                    ).first()

                if paciente:
                    if paciente.nombre_hospital != hospital or paciente.estado_paciente != estado:
                        mov = models.PacienteMovimiento(
                            paciente_id=paciente.id,
                            hospital_anterior=paciente.nombre_hospital,
                            hospital_nuevo=hospital,
                            estado_anterior=paciente.estado_paciente,
                            estado_nuevo=estado,
                            batch_date_anterior=paciente.batch_date,
                            batch_date_nuevo=batch_date_str,
                        )
                        db.add(mov)
                        movimientos += 1

                    paciente.nombre_hospital = hospital
                    paciente.estado_paciente = estado
                    paciente.batch_date = batch_date_str
                    if cedula and not paciente.cedula:
                        paciente.cedula = cedula
                    if variantes:
                        paciente.nombre_variantes = variantes
                    if telefono:
                        paciente.telefono = telefono
                    if observaciones:
                        paciente.observaciones = observaciones
                    paciente.fuente = "excel"
                    actualizados += 1
                else:
                    nuevo = models.PacienteHospitalizado(
                        nombres_apellidos=nombre_principal,
                        nombre_variantes=variantes,
                        nombre_hospital=hospital,
                        edad=edad,
                        cedula=cedula,
                        telefono=telefono,
                        procedencia=direccion,
                        observaciones=observaciones,
                        estado_paciente=estado,
                        batch_date=batch_date_str,
                        fuente="excel",
                    )
                    db.add(nuevo)
                    nuevos += 1

                if i % 100 == 0:
                    db.commit()
                    if total > 0:
                        _jobs[job_id]["progress"] = int((i / total) * 100)

            except Exception as exc:
                errores += 1
                errores_detalle.append(f"Fila {row_num}: {exc}")

        db.commit()
        _jobs[job_id].update({
            "status": "done",
            "progress": 100,
            "result": {
                "importados": nuevos,
                "actualizados": actualizados,
                "movimientos_detectados": movimientos,
                "errores": errores,
                "total_procesadas": total,
                "batch_date": batch_date_str,
            },
        })

    except Exception as exc:
        _jobs[job_id].update({"status": "error", "error": str(exc)})

    finally:
        db.close()


@router.post("/import/excel-pacientes")
async def import_excel_pacientes(
    request: Request,
    file: UploadFile = File(...),
    batch_date_str: str = Form(...),
    x_csrf_token: str = Header(""),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")

    fname = (file.filename or "").lower()
    if not (fname.endswith(".xlsx") or fname.endswith(".xls")):
        raise HTTPException(status_code=400, detail="El archivo debe ser .xlsx o .xls")

    file_bytes = await read_limited(file)
    if file_bytes is None:
        raise HTTPException(status_code=413, detail="Archivo demasiado grande (máx 10MB)")

    job_id = str(uuid.uuid4())[:8]
    _jobs[job_id] = {
        "status": "pending",
        "progress": 0,
        "result": None,
        "error": None,
        "created_at": datetime.now(),
    }

    threading.Thread(
        target=_procesar_excel,
        args=(job_id, file_bytes, batch_date_str, admin),
        daemon=True,
    ).start()

    return {"job_id": job_id, "status": "pending"}


@router.get("/import/jobs/{job_id}")
async def get_import_job(job_id: str, request: Request):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")

    job = _jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job no encontrado")

    return {
        "status": job["status"],
        "progress": job["progress"],
        "result": job["result"],
        "error": job["error"],
    }


# --- Centros de acopio ---

@router.post("/acopio/centros/nuevo")
async def nuevo_centro_acopio(
    request: Request,
    nombre: str = Form(...),
    zona: Optional[str] = Form(None),
    direccion: Optional[str] = Form(None),
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")
    centro = models.CentroAcopio(
        nombre=nombre,
        zona=zona or None,
        direccion=direccion or None,
    )
    db.add(centro)
    db.commit()
    db.refresh(centro)
    return {"ok": True, "id": centro.id, "nombre": centro.nombre}


@router.delete("/acopio/centros/{centro_id}")
async def desactivar_centro_acopio(
    centro_id: int,
    request: Request,
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")
    centro = db.query(models.CentroAcopio).filter(
        models.CentroAcopio.id == centro_id
    ).first()
    if not centro:
        raise HTTPException(status_code=404, detail="Centro no encontrado")
    centro.activo = False
    db.commit()
    return {"ok": True}


@router.delete("/acopio/reportes/{reporte_id}")
async def eliminar_acopio_reporte(
    reporte_id: int,
    request: Request,
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")
    reporte = db.query(models.AcopioReporte).filter(
        models.AcopioReporte.id == reporte_id
    ).first()
    if not reporte:
        raise HTTPException(status_code=404, detail="Reporte no encontrado")
    db.delete(reporte)
    db.commit()
    return {"ok": True}


# --- Solicitudes de centros de acopio ---

@router.get("/acopio/solicitudes")
async def listar_solicitudes_acopio(
    request: Request,
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    solicitudes = db.query(models.CentroAcopioSolicitud).filter(
        models.CentroAcopioSolicitud.estado == "pendiente"
    ).order_by(sqldesc(models.CentroAcopioSolicitud.created_at)).all()
    return [_schemas.CentroSolicitudResponse.model_validate(s) for s in solicitudes]


@router.post("/acopio/solicitudes/{solicitud_id}/aprobar")
async def aprobar_solicitud_acopio(
    solicitud_id: int,
    request: Request,
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")
    solicitud = db.query(models.CentroAcopioSolicitud).filter(
        models.CentroAcopioSolicitud.id == solicitud_id
    ).first()
    if not solicitud:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    centro = models.CentroAcopio(
        nombre=solicitud.nombre,
        zona=solicitud.zona or None,
        direccion=solicitud.direccion or None,
    )
    db.add(centro)
    solicitud.estado = "aprobado"
    db.commit()
    db.refresh(centro)
    return {"ok": True, "centro_id": centro.id}


@router.post("/acopio/solicitudes/{solicitud_id}/rechazar")
async def rechazar_solicitud_acopio(
    solicitud_id: int,
    request: Request,
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")
    solicitud = db.query(models.CentroAcopioSolicitud).filter(
        models.CentroAcopioSolicitud.id == solicitud_id
    ).first()
    if not solicitud:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    solicitud.estado = "rechazado"
    db.commit()
    return {"ok": True}


# --- Notificaciones de personas localizadas ---

@router.get("/notificaciones/localizado")
async def listar_notificaciones_localizado(
    request: Request,
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")

    notificaciones = db.query(models.PersonaNotificacionLocalizado).filter(
        models.PersonaNotificacionLocalizado.estado == "pendiente"
    ).order_by(sqldesc(models.PersonaNotificacionLocalizado.created_at)).all()

    result = []
    for n in notificaciones:
        persona = db.query(models.PersonaDesaparecida).filter(
            models.PersonaDesaparecida.id == n.persona_id
        ).first()
        result.append({
            "notificacion": _schemas.NotificacionLocalizadoResponse.model_validate(n),
            "persona": {
                "id": persona.id,
                "nombres_apellidos": persona.nombres_apellidos,
                "cedula": persona.cedula,
                "estado": persona.estado,
            } if persona else None,
        })
    return result


@router.post("/notificaciones/{notificacion_id}/confirmar")
async def confirmar_notificacion_localizado(
    notificacion_id: int,
    request: Request,
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")

    notificacion = db.query(models.PersonaNotificacionLocalizado).filter(
        models.PersonaNotificacionLocalizado.id == notificacion_id
    ).first()
    if not notificacion:
        raise HTTPException(status_code=404, detail="No encontrado")

    notificacion.estado = "confirmado"
    persona = db.query(models.PersonaDesaparecida).filter(
        models.PersonaDesaparecida.id == notificacion.persona_id
    ).first()
    if persona:
        persona.estado = "localizado"
    db.commit()
    return {"ok": True}


@router.post("/notificaciones/{notificacion_id}/rechazar")
async def rechazar_notificacion_localizado(
    notificacion_id: int,
    request: Request,
    x_csrf_token: str = Header(""),
    db: Session = Depends(get_db),
):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(status_code=401, detail="No autorizado")
    if not validate_csrf(request, x_csrf_token, admin):
        raise HTTPException(status_code=403, detail="CSRF token inválido")

    notificacion = db.query(models.PersonaNotificacionLocalizado).filter(
        models.PersonaNotificacionLocalizado.id == notificacion_id
    ).first()
    if not notificacion:
        raise HTTPException(status_code=404, detail="No encontrado")

    notificacion.estado = "rechazado"
    db.commit()
    return {"ok": True}
